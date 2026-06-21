"""Ferramentas basicas de DOCX (leitura e criacao simples de documentos)."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from .logger import get_logger

logger = get_logger("docx_tools")


def ler_docx(caminho: Path) -> str:
    documento = Document(str(caminho))
    texto = "\n".join(paragrafo.text for paragrafo in documento.paragraphs)
    logger.info(f"DOCX lido: {caminho} ({len(texto)} caracteres).")
    return texto


def criar_docx_simples(caminho: Path, titulo: str, paragrafos: list[str]) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    documento = Document()
    documento.add_heading(titulo, level=1)
    for paragrafo in paragrafos:
        documento.add_paragraph(paragrafo)
    documento.save(str(caminho))
    logger.info(f"DOCX criado em {caminho}.")
    return caminho
